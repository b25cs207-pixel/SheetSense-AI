import io
import re
import pandas as pd
from docx import Document
from backend.parsers.base import ParsedData
from backend.parsers.document_structurer import structure_text_content

def parse_docx(file_bytes: bytes, filename: str = "upload.docx") -> ParsedData:
    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        full_text = "\n".join(paragraphs)
        
        # 1. Check embedded Word tables
        if doc.tables:
            for table in doc.tables:
                rows_data = []
                for row in table.rows:
                    row_vals = [cell.text.strip() for cell in row.cells]
                    if any(row_vals):
                        rows_data.append(row_vals)
                if len(rows_data) >= 2:
                    headers = rows_data[0]
                    unique_headers = []
                    seen = set()
                    for idx, h in enumerate(headers):
                        name = h if h and h not in seen else f"Col_{idx+1}"
                        seen.add(name)
                        unique_headers.append(name)
                    df = pd.DataFrame(rows_data[1:], columns=unique_headers)
                    for c in df.columns:
                        df[c] = pd.to_numeric(df[c], errors="ignore")
                    return ParsedData(
                        df=df,
                        text_content=full_text[:3000],
                        metadata={"tables_found": len(doc.tables), "rows": len(df), "cols": len(df.columns)},
                        is_structured=True
                    )

        # 2. Resilient fallback: Structure extracted text
        df_struct, struct_warnings, struct_meta = structure_text_content(full_text, filename)
        if df_struct is not None:
            return ParsedData(
                df=df_struct,
                text_content=full_text,
                metadata=struct_meta,
                warnings=struct_warnings,
                is_structured=True
            )

        return ParsedData(
            df=None,
            text_content=full_text,
            metadata={"paragraph_count": len(paragraphs)},
            warnings=["Document text extracted, but no quantitative records could be inferred."],
            is_structured=False
        )
    except Exception as e:
        return ParsedData(
            df=None,
            warnings=[f"Failed to parse DOCX: {str(e)}"],
            is_structured=False
        )

def parse_txt(file_bytes: bytes, filename: str = "upload.txt") -> ParsedData:
    try:
        text = file_bytes.decode("utf-8-sig", errors="replace")
        lines = [l.strip() for l in text.split("\n") if l.strip()]
        
        # 1. Delimiter detection
        if len(lines) >= 3:
            first = lines[0]
            delims = [",", "\t", "|", ";"]
            counts = {d: first.count(d) for d in delims}
            best_d = max(counts, key=counts.get)
            if counts[best_d] >= 1:
                try:
                    df = pd.read_csv(io.StringIO(text), sep=best_d, on_bad_lines="skip")
                    if len(df.columns) >= 2 and len(df) >= 2:
                        return ParsedData(
                            df=df,
                            metadata={"detected_delimiter": best_d, "rows": len(df), "cols": len(df.columns)},
                            is_structured=True
                        )
                except Exception:
                    pass

        # 2. Resilient fallback
        df_struct, struct_warnings, struct_meta = structure_text_content(text, filename)
        if df_struct is not None:
            return ParsedData(
                df=df_struct,
                text_content=text,
                metadata=struct_meta,
                warnings=struct_warnings,
                is_structured=True
            )

        return ParsedData(
            df=None,
            text_content=text,
            metadata={"line_count": len(lines)},
            warnings=["Plain text document processed."],
            is_structured=False
        )
    except Exception as e:
        return ParsedData(
            df=None,
            warnings=[f"Failed to read TXT: {str(e)}"],
            is_structured=False
        )
